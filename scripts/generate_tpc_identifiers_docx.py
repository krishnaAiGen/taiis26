"""
Generate TAIIS 2026 TPC member Word list with ORCID / Scopus / Google Scholar IDs.
Data source: src/config/siteConfig.ts (Technical Program Committee).
Lookups: OpenAlex + ORCID public API (best-effort affiliation matching).
"""

from __future__ import annotations

import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SITE_CONFIG = ROOT / "src" / "config" / "siteConfig.ts"
OUT_DOCX = ROOT / "docs" / "TAIIS2026_TPC_Member_Identifiers.docx"
CACHE_FILE = ROOT / "scripts" / ".tpc_identifier_cache.json"

MAILTO = "taiis2026@asia.edu.tw"
REQUEST_DELAY_SEC = 0.35

STOPWORDS = {
    "university",
    "institute",
    "college",
    "national",
    "technology",
    "india",
    "china",
    "taiwan",
    "kingdom",
    "saudi",
    "arabia",
    "republic",
    "states",
    "united",
    "the",
    "of",
    "and",
    "uk",
    "usa",
    "uae",
    "ksa",
}


@dataclass
class MemberIds:
    name: str
    affiliation: str
    orcid: str = ""
    scopus_id: str = ""
    google_scholar: str = ""
    match_note: str = ""
    openalex_id: str = ""


def http_get_json(url: str, headers: dict | None = None) -> dict | list | None:
    hdrs = {"User-Agent": f"TAIIS2026-site-script (mailto:{MAILTO})"}
    if headers:
        hdrs.update(headers)
    req = urllib.request.Request(url, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=45) as resp:
            return json.load(resp)
    except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, json.JSONDecodeError):
        return None


def parse_tpc_members() -> list[tuple[str, str]]:
    text = SITE_CONFIG.read_text(encoding="utf-8")
    tpc_match = re.search(
        r'role:\s*"Technical Program Committee",\s*members:\s*\[(.*?)\],\s*\n\s*\},',
        text,
        re.DOTALL,
    )
    if not tpc_match:
        raise SystemExit("Could not parse TPC block from siteConfig.ts")
    block = tpc_match.group(1)
    members: list[tuple[str, str]] = []
    for m in re.finditer(
        r'name:\s*"([^"]+)"[\s\S]*?affiliation:\s*"([^"]+)"',
        block,
    ):
        members.append((m.group(1).strip(), m.group(2).strip()))
    return members


def affiliation_keywords(affiliation: str) -> set[str]:
    tokens = re.split(r"[\s,./()\-–]+", affiliation.lower())
    return {t for t in tokens if len(t) >= 4 and t not in STOPWORDS}


def institution_names(author: dict) -> list[str]:
    names: list[str] = []
    for aff in author.get("affiliations") or []:
        inst = aff.get("institution") or {}
        if inst.get("display_name"):
            names.append(inst["display_name"].lower())
    for inst in author.get("last_known_institutions") or []:
        if inst.get("display_name"):
            names.append(inst["display_name"].lower())
    return names


def score_author(affiliation: str, author: dict) -> int:
    aff_low = affiliation.lower()
    keys = affiliation_keywords(affiliation)
    score = 0
    display = (author.get("display_name") or "").lower()
    for key in keys:
        if key in aff_low:
            for inst in institution_names(author):
                if key in inst:
                    score += 4
        if key in display:
            score += 1
    return score


def extract_scopus_from_ids(ids: dict) -> str:
    if not ids:
        return ""
    scopus = ids.get("scopus") or ids.get("Scopus") or ""
    if scopus:
        m = re.search(r"authorId=(\d+)", scopus)
        if m:
            return m.group(1)
        m = re.search(r"(\d{8,12})", scopus)
        if m:
            return m.group(1)
    return ""


def fetch_orcid_extras(orcid_url: str) -> tuple[str, str]:
    """Return (scopus_id, google_scholar_url) from ORCID record."""
    if not orcid_url:
        return "", ""
    orcid = orcid_url.rstrip("/").split("/")[-1]
    record = http_get_json(
        f"https://pub.orcid.org/v3.0/{orcid}/record",
        headers={"Accept": "application/json"},
    )
    if not isinstance(record, dict):
        return "", ""

    scopus = ""
    scholar = ""

    ext_block = (record.get("person") or {}).get("external-identifiers") or {}
    for item in ext_block.get("external-identifier") or []:
        id_type = (item.get("external-id-type") or "").lower()
        value = item.get("external-id-value") or ""
        url_obj = item.get("external-id-url") or {}
        url = url_obj.get("value") if isinstance(url_obj, dict) else ""
        if "scopus" in id_type and value:
            scopus = re.sub(r"\D", "", value) or value
        if "scopus" in (url or "").lower():
            m = re.search(r"authorId=(\d+)", url)
            if m:
                scopus = m.group(1)

    url_block = (record.get("person") or {}).get("researcher-urls") or {}
    for item in url_block.get("researcher-url") or []:
        url_obj = item.get("url") or {}
        url = (url_obj.get("value") or "").strip()
        if "scholar.google" in url.lower():
            scholar = url

    return scopus, scholar


def lookup_member(name: str, affiliation: str) -> MemberIds:
    result = MemberIds(name=name, affiliation=affiliation)
    query = urllib.parse.quote(name)
    url = (
        f"https://api.openalex.org/authors?search={query}&per_page=8"
        f"&mailto={urllib.parse.quote(MAILTO)}"
    )
    data = http_get_json(url)
    time.sleep(REQUEST_DELAY_SEC)

    if not isinstance(data, dict) or not data.get("results"):
        result.match_note = "No OpenAlex match"
        return result

    ranked = sorted(
        data["results"],
        key=lambda a: score_author(affiliation, a),
        reverse=True,
    )
    best = ranked[0]
    best_score = score_author(affiliation, best)
    second_score = score_author(affiliation, ranked[1]) if len(ranked) > 1 else -1

    if best_score == 0:
        result.match_note = "Low confidence (name only)"
    elif second_score == best_score:
        result.match_note = "Ambiguous affiliation match"
    else:
        result.match_note = "Affiliation matched" if best_score >= 4 else "Partial affiliation match"

    ids = best.get("ids") or {}
    orcid = ids.get("orcid") or best.get("orcid") or ""
    if orcid and not orcid.startswith("http"):
        orcid = f"https://orcid.org/{orcid}"
    result.orcid = orcid.replace("https://orcid.org/", "") if orcid else ""
    result.openalex_id = (ids.get("openalex") or best.get("id") or "").replace(
        "https://openalex.org/", ""
    )
    result.scopus_id = extract_scopus_from_ids(ids)

    if orcid:
        scopus_extra, scholar = fetch_orcid_extras(orcid)
        time.sleep(REQUEST_DELAY_SEC)
        if scopus_extra and not result.scopus_id:
            result.scopus_id = scopus_extra
        result.google_scholar = scholar

    # Full author record may include Scopus in merged IDs
    openalex_key = result.openalex_id
    if openalex_key and not result.scopus_id:
        full = http_get_json(
            f"https://api.openalex.org/authors/{openalex_key}?mailto={urllib.parse.quote(MAILTO)}"
        )
        time.sleep(REQUEST_DELAY_SEC)
        if isinstance(full, dict):
            result.scopus_id = extract_scopus_from_ids(full.get("ids") or {}) or result.scopus_id

    return result


def load_cache() -> dict[str, dict]:
    if CACHE_FILE.is_file():
        return json.loads(CACHE_FILE.read_text(encoding="utf-8"))
    return {}


def save_cache(cache: dict[str, dict]) -> None:
    CACHE_FILE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")


def build_docx(rows: list[MemberIds]) -> None:
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = doc.add_heading("TAIIS 2026 — Technical Program Committee", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph(
        "Member identifiers (ORCID, Scopus Author ID, Google Scholar profile). "
        "Values were retrieved automatically from OpenAlex and ORCID public records "
        "using name and affiliation matching. Please verify low-confidence or blank entries."
    )
    intro.paragraph_format.space_after = Pt(12)

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = [
        "#",
        "Name",
        "Affiliation",
        "ORCID",
        "Scopus Author ID",
        "Google Scholar URL",
        "Lookup note",
    ]
    for i, label in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = label
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True

    for idx, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        cells[0].text = str(idx)
        cells[1].text = row.name
        cells[2].text = row.affiliation
        cells[3].text = (
            f"https://orcid.org/{row.orcid}" if row.orcid else ""
        )
        cells[4].text = (
            f"https://www.scopus.com/authors/detail.uri?authorId={row.scopus_id}"
            if row.scopus_id
            else ""
        )
        cells[5].text = row.google_scholar
        cells[6].text = row.match_note

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)

    doc.add_paragraph()
    foot = doc.add_paragraph(
        f"Generated from {SITE_CONFIG.name}. Total TPC members: {len(rows)}."
    )
    foot.runs[0].font.size = Pt(9)

    section = doc.sections[0]
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    doc.save(OUT_DOCX)


def main() -> None:
    members = parse_tpc_members()
    cache = load_cache()
    rows: list[MemberIds] = []

    for name, affiliation in members:
        key = name.lower()
        if key in cache:
            rows.append(MemberIds(**cache[key]))
            continue
        info = lookup_member(name, affiliation)
        rows.append(info)
        cache[key] = info.__dict__
        save_cache(cache)
        line = f"  {name}: ORCID={info.orcid or '-'} Scopus={info.scopus_id or '-'}"
        print(line.encode("ascii", errors="replace").decode("ascii"))

    build_docx(rows)
    print(f"\nWrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
