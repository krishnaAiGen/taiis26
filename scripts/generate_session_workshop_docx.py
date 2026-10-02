"""
Generate Word documents for TAIIS 2026 special sessions and workshops.
Output: docs/special-sessions/*.docx, docs/workshops/*.docx
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "src" / "config"
OUT_SS = ROOT / "docs" / "special-sessions"
OUT_WS = ROOT / "docs" / "workshops"


def slugify(text: str, max_len: int = 48) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_")
    return s[:max_len].rstrip("_")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_document(
    *,
    event_type: str,
    code: str,
    title: str,
    acronym: str | None,
    aim: list[str],
    topics: list[str],
    section_titles: dict[str, str],
    committees_note: str,
    submission_dates: list[dict[str, str]],
) -> Document:
    doc = Document()
    title_p = doc.add_heading("TAIIS 2026", level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"{event_type}: {code}")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph(title, style="Intense Quote").alignment = WD_ALIGN_PARAGRAPH.CENTER
    if acronym:
        ac = doc.add_paragraph(f"({acronym})")
        ac.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    add_heading(doc, section_titles["aimAndScope"], level=1)
    for para in aim:
        doc.add_paragraph(para)

    add_heading(doc, section_titles["topics"], level=1)
    add_bullets(doc, topics)

    add_heading(doc, section_titles["committees"], level=1)
    doc.add_paragraph(committees_note)

    add_heading(doc, section_titles["submissionDates"], level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Milestone"
    hdr[1].text = "Date"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in submission_dates:
        cells = table.add_row().cells
        cells[0].text = row["label"]
        cells[1].text = row["date"]

    doc.add_paragraph()
    foot = doc.add_paragraph(
        "International Conference on Trustworthy AI and Intelligent IoT Systems (TAIIS 2026) · Asia University, Taichung, Taiwan"
    )
    foot.runs[0].font.size = Pt(9)

    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    return doc


def main() -> None:
    shared = json.loads((CONFIG / "sessionWorkshopShared.json").read_text(encoding="utf-8"))
    ss_data = json.loads((CONFIG / "specialSessions.json").read_text(encoding="utf-8"))
    ws_data = json.loads((CONFIG / "workshops.json").read_text(encoding="utf-8"))

    OUT_SS.mkdir(parents=True, exist_ok=True)
    OUT_WS.mkdir(parents=True, exist_ok=True)

    for session in ss_data["sessions"]:
        doc = build_document(
            event_type="Special Session",
            code=session["code"],
            title=session["title"],
            acronym=None,
            aim=session["aimAndScope"],
            topics=session["topics"],
            section_titles=ss_data["sectionTitles"],
            committees_note=shared["committeesNote"],
            submission_dates=shared["submissionDates"],
        )
        fname = f"{session['code']}_{slugify(session['title'])}.docx"
        path = OUT_SS / fname
        doc.save(path)
        print("Wrote", path.relative_to(ROOT))

    for workshop in ws_data["workshops"]:
        doc = build_document(
            event_type="Workshop",
            code=workshop["code"],
            title=workshop["title"],
            acronym=workshop.get("acronym"),
            aim=workshop["aimAndScope"],
            topics=workshop["topics"],
            section_titles=ws_data["sectionTitles"],
            committees_note=shared["committeesNote"],
            submission_dates=shared["submissionDates"],
        )
        fname = f"{workshop['code']}_{slugify(workshop.get('acronym') or workshop['title'])}.docx"
        path = OUT_WS / fname
        doc.save(path)
        print("Wrote", path.relative_to(ROOT))


if __name__ == "__main__":
    main()
